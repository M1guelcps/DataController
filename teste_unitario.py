import openpyxl
from io import BytesIO
from pathlib import Path

from app.core.database import SessionLocal
from app.repositories.cliente_repository import ClienteRepository
from core.config import settings
from app.domain.cliente_model import Cliente

# ============================================================
# 1. CONFIGURAÇÃO DOS ARQUIVOS
# ============================================================

BASE_DIR = Path(__file__).resolve().parent

caminho_template = BASE_DIR / "assets" / "template_clientes.xlsx"
caminho_saida = BASE_DIR / "assets" / "teste_resultado.xlsx"


# ============================================================
# 2. VERIFICA SE O TEMPLATE EXISTE
# ============================================================

print(f"Template: {caminho_template}")

if not caminho_template.exists():
    print("ERRO: template não encontrado!")
    exit()

print("Template encontrado!")


# ============================================================
# 3. ABRE O TEMPLATE
# ============================================================

workbook = openpyxl.load_workbook(caminho_template)

print("Abas encontradas:")
print(workbook.sheetnames)

planilha = workbook["clientes"]


# ============================================================
# 4. ABRE A CONEXÃO COM O BANCO
# ============================================================

db = SessionLocal()
quantidade = db.query(Cliente).count()

print(f"Quantidade de registros na tabela clientes: {quantidade}")

print()
print("======================================")
print("CONFIGURAÇÃO DO BANCO")
print("======================================")

print(f"DATABASE_URL: {settings.DATABASE_URL}")




try:

    # Busca os clientes através do Repository
    clientes = ClienteRepository.buscar_todos(db)

    print()
    print("======================================")
    print("CLIENTES ENCONTRADOS NO BANCO")
    print("======================================")

    print(f"Quantidade de clientes: {len(clientes)}")
    print()


    # ========================================================
    # 5. MOSTRA OS CLIENTES ENCONTRADOS
    # ========================================================

    for cliente in clientes:
        print(
            f"ID: {cliente.id} | "
            f"Nome: {cliente.nome} | "
            f"CNPJ: {cliente.cnpj} | "
            f"Email: {cliente.email}"
        )


    # ========================================================
    # 6. PREENCHE O EXCEL
    # ========================================================

    print()
    print("======================================")
    print("PREENCHENDO EXCEL")
    print("======================================")

    linha_atual = 8

    for cliente in clientes:

        print(
            f"Preenchendo linha {linha_atual}: "
            f"{cliente.nome}"
        )

        planilha.cell(
            row=linha_atual,
            column=1,
            value=cliente.nome
        )

        planilha.cell(
            row=linha_atual,
            column=2,
            value=cliente.email
        )

        planilha.cell(
            row=linha_atual,
            column=3,
            value=cliente.cnpj
        )

        planilha.cell(
            row=linha_atual,
            column=4,
            value=cliente.id
        )

        linha_atual += 1


    # ========================================================
    # 7. SALVA O EXCEL EM MEMÓRIA
    # ========================================================

    arquivo_em_memoria = BytesIO()

    workbook.save(arquivo_em_memoria)

    arquivo_em_memoria.seek(0)


    # ========================================================
    # 8. SALVA UMA CÓPIA PARA PODERMOS TESTAR
    # ========================================================

    with open(caminho_saida, "wb") as arquivo:
        arquivo.write(arquivo_em_memoria.getvalue())


    print()
    print("======================================")
    print("TESTE CONCLUÍDO")
    print("======================================")

    print(f"Arquivo gerado em:")
    print(caminho_saida)


finally:

    # Fecha a conexão com o banco
    db.close()